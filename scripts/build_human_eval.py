#!/usr/bin/env python3
"""Human-evaluation set: the same 150 original source sentences (3 per drug, seed 42) with every triple each model
extracted from the chunk that contains the sentence.

Model A = EDC only (drug-kg/output/experiments/text_edc; triples from its explorer HTML, the only copy on this server)
Model B = pronoun resolution + sentence decomposition + EDC (drug-kg/output/experiments/coref_decomp_edc_r2)

EDC extracts from chunks of up to 1,100 characters (about 8 original sentences) and records the chunk, not the
sentence, a triple came from. Links used here are therefore exact but chunk-level, and nothing is inferred:
  original sentence -> A chunk : by re-running the chunker's own grouping rule on the original text file (line-level)
  original sentence -> B chunk : by the original-sentence record ids stored in B's chunks
  triple -> chunk              : A: the edge's chunk id (first provenance only, all that the HTML keeps);
                                 B: every chunk id in the triple's provenance
Each triple row carries "From this source? (Yes/No)" for the evaluator to drop triples of the chunk's other sentences.
No similarity matching, no triple editing, no re-extraction.
"""

from __future__ import annotations

import argparse
import collections
import csv
import json
import random
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "drug-kg"))
from classify_drugs_pilot_sentences import ensure_nltk, extract_directory_records  # noqa: E402
from drugkg.chunk import _split_long  # noqa: E402

EXCLUDE_SECTIONS = {"related-drugs", "faq", "references", "more-about", "further-information", "images", "popular-faq"}
MAX_CHARS, MIN_CHARS = 1100, 150   # drug-kg/config.yaml chunking (identical in both experiments)


def slugify(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


def norm(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def source_id(source_file: str, line: int, idx: int) -> str:
    return f"{Path(source_file).stem}:L{line}:S{idx}"


# --------------------------------------------------------------------------- model A: line -> chunk
def a_line_chunks(path: Path) -> dict[int, list[tuple[str, str]]]:
    """Replicates drugkg.inputs.text_file_to_record + drugkg.chunk.chunk_record, keeping line numbers.
    Returns line number -> [(chunk_id, piece text)] (several pieces only for lines longer than 1,100 chars)."""
    slug = path.stem
    sections, cur = [], None
    for no, line in enumerate(path.read_text(encoding="utf-8").split("\n"), start=1):
        if line.startswith("# ") and not sections and cur is None:
            continue
        if line.startswith("Source: "):
            continue
        if line.startswith("## "):
            cur = {"heading": line[3:].strip(), "lines": []}
            sections.append(cur)
        elif line.strip() and cur is not None:
            cur["lines"].append((no, line.rstrip()))
    out: dict[int, list[tuple[str, str]]] = collections.defaultdict(list)
    seen: dict[str, int] = {}
    for sec in sections:
        sid = slugify(sec["heading"])
        seen[sid] = seen.get(sid, 0) + 1
        if seen[sid] > 1:
            sid = f"{sid}-{seen[sid]}"
        if sid.split(":", 1)[-1] in EXCLUDE_SECTIONS or sid in EXCLUDE_SECTIONS:
            continue
        pieces = [(no, p) for no, line in sec["lines"] for p in _split_long(line, MAX_CHARS)]
        groups, g, n = [], [], 0
        for no, p in pieces:
            if g and n + len(p) + 1 > MAX_CHARS:
                groups.append(g)
                g, n = [], 0
            g.append((no, p))
            n += len(p) + 1
        if g:
            groups.append(g)
        if len(groups) >= 2 and sum(len(p) for _, p in groups[-1]) < MIN_CHARS:
            groups[-2].extend(groups.pop())
        for i, grp in enumerate(groups):
            if len("\n".join(p for _, p in grp).strip()) < 20:
                continue
            for no, p in grp:
                out[no].append((f"{slug}#{sid}#{i}", p))
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--raw-dir", type=Path, default=ROOT / "data/text-only/text-only")
    ap.add_argument("--a-html", type=Path, default=ROOT / "drug-kg/output/viz/text_edc/kg_explorer.html")
    ap.add_argument("--b-exp", type=Path, default=ROOT / "drug-kg/output/experiments/coref_decomp_edc_r2")
    ap.add_argument("--trace", type=Path, default=ROOT / "data/text-only/coref_pipeline/4_trace/trace.jsonl")
    ap.add_argument("--out", type=Path, default=ROOT / "data/human_eval")
    ap.add_argument("--per-drug", type=int, default=3)
    ap.add_argument("--seed", type=int, default=42)
    args = ap.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)

    # ---- original sentences (same splitter as the CoDe-KG classification step)
    ensure_nltk()
    raw, _ = extract_directory_records(args.raw_dir)

    # ---- model A: chunks + triples
    html = args.a_html.read_text(encoding="utf-8")
    A = json.loads(re.search(r'<script\b[^>]*\bid=["\']kg-data["\'][^>]*>(.*?)</script>', html, re.S).group(1))
    a_label = {n["id"]: n["l"] for n in A["nodes"]}
    a_triples = collections.defaultdict(list)
    for i, e in enumerate(A["edges"]):
        if "edc" in e["src"] and e.get("ch"):
            a_triples[e["ch"]].append({"triple_id": f"A{i:05d}", "subject": a_label[e["s"]], "relation": e["r"],
                                       "object": a_label[e["o"]], "count_in_kg": e["c"]})
    a_lines = {p.name: a_line_chunks(p) for p in args.raw_dir.glob("*.txt")}
    rebuilt = collections.defaultdict(list)   # chunk id -> pieces in file order (line numbers ascend)
    for f in a_lines.values():
        for no in sorted(f):
            for cid, piece in f[no]:
                rebuilt[cid].append(piece)
    # the line-level rebuild must reproduce every chunk text the EDC-only run used, character for character
    mismatch = [cid for cid, text in A["chunks"].items() if "\n".join(rebuilt.get(cid, [])).strip() != text]
    if mismatch:
        print(json.dumps({"error": "EDC-only chunk rebuild differs from the original chunks", "chunks": mismatch[:10]}))
        return 1
    a_chunk_text = A["chunks"]

    # ---- model B: chunks (with original-sentence record ids) + triples
    trace = {}
    for line in args.trace.open(encoding="utf-8"):
        t = json.loads(line)
        f, ln, si, _ = t["record_id"].split("|")
        trace[(f, int(ln), int(si))] = t
    b_rec_chunks = collections.defaultdict(set)
    b_chunk_text = {}
    for path in (args.b_exp / "chunks").glob("*.jsonl"):
        for line in path.open(encoding="utf-8"):
            c = json.loads(line)
            b_chunk_text[c["chunk_id"]] = c["text"]
            for s in c["sentences"]:
                b_rec_chunks[s["record_id"]].add(c["chunk_id"])
    b_triples = collections.defaultdict(list)
    for k, line in enumerate((args.b_exp / "kg/triples.jsonl").open(encoding="utf-8")):
        t = json.loads(line)
        chunks = [p["chunk_id"] for p in t.get("provenance", []) if p.get("source") == "edc" and p.get("chunk_id")]
        for j, cid in enumerate(dict.fromkeys(chunks)):
            b_triples[cid].append({"triple_id": f"B{k:05d}", "subject": t["subject"], "relation": t["relation"],
                                   "object": t["object"], "first_provenance": j == 0, "count_in_kg": t["count"]})

    # ---- link every original sentence to both models; keep only exactly linkable ones
    universe = collections.defaultdict(list)
    reasons = collections.Counter()
    for r in raw:
        key = (r["source_file"], r["line_number"], r["sentence_index"])
        a_ch = [cid for cid, piece in a_lines[r["source_file"]].get(r["line_number"], [])
                if norm(r["sentence"]) in norm(piece[2:] if piece.startswith("- ") else piece)]
        t = trace.get(key)
        if not a_ch:
            reasons["not in any EDC-only chunk (excluded section or split across pieces)"] += 1
            continue
        if t is None:
            reasons["no record in the decomposition pipeline"] += 1
            continue
        if t["raw_match"] != "sentence":
            reasons["pronoun resolution changed how this line splits into sentences (no exact sentence id)"] += 1
            continue
        b_ch = sorted(b_rec_chunks.get(t["record_id"], set()))
        if not b_ch:
            reasons["not in any preprocessing+EDC chunk (excluded heading)"] += 1
            continue
        universe[r["source_file"]].append({
            "source_id": source_id(r["source_file"], r["line_number"], r["sentence_index"]), "drug": r["drug"],
            "source_file": r["source_file"], "section": r["section_title"], "line": r["line_number"],
            "original": r["sentence"], "coref": t["coref_sentence"], "coref_changed": bool(t["coref_edits"]),
            "label": t["label"], "status": t["status"],
            "decomposed": [{"id": f"{source_id(r['source_file'], r['line_number'], r['sentence_index'])}:D{d['output_index'] + 1}",
                            "sentence": d["sentence"]} for d in t["decomposed"]],
            "a_chunks": sorted(set(a_ch)), "b_chunks": b_ch})

    files = sorted(p.name for p in args.raw_dir.glob("*.txt"))
    short = {f: len(universe[f]) for f in files if len(universe[f]) < args.per_drug}
    if short:
        print(json.dumps({"error": "fewer linkable sentences than per-drug quota", "drugs": short}), flush=True)
        return 1

    rng = random.Random(args.seed)
    sampled = []
    for f in files:
        for s in rng.sample(universe[f], args.per_drug):
            sampled.append(s)
    for k, s in enumerate(sampled, 1):
        s["no"] = k

    # ---- rows
    a_rows, b_rows = [], []
    for s in sampled:
        seen = set()
        for cid in s["a_chunks"]:
            for t in a_triples.get(cid, []):
                if t["triple_id"] not in seen:
                    seen.add(t["triple_id"])
                    a_rows.append({**s, "chunk_id": cid, **t})
        if not seen:
            a_rows.append({**s, "chunk_id": ";".join(s["a_chunks"]), "triple_id": "", "subject": "", "relation": "",
                           "object": "(no triple extracted from this chunk)", "count_in_kg": ""})
        seen = set()
        for cid in s["b_chunks"]:
            for t in b_triples.get(cid, []):
                if t["triple_id"] not in seen:
                    seen.add(t["triple_id"])
                    b_rows.append({**s, "chunk_id": cid, **t})
        if not seen:
            b_rows.append({**s, "chunk_id": ";".join(s["b_chunks"]), "triple_id": "", "subject": "", "relation": "",
                           "object": "(no triple extracted from this chunk)", "first_provenance": "", "count_in_kg": ""})

    def dec_text(s):
        return "\n".join(f"[{d['id']}] {d['sentence']}" for d in s["decomposed"])

    def dec_ids(s):
        return "; ".join(d["id"] for d in s["decomposed"])

    with (args.out / "sampled_sources.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["No", "Source ID", "Drug name", "Original source sentence", "Section", "Line",
                    "Coreference-resolved sentence", "Pronoun changed", "Sentence type (BERT)", "Decomposition status",
                    "Decomposed sentences", "Decomposed sentence IDs", "EDC-only chunk ID", "Preprocessing+EDC chunk ID(s)"])
        for s in sampled:
            w.writerow([s["no"], s["source_id"], s["drug"], s["original"], s["section"], s["line"], s["coref"],
                        "Yes" if s["coref_changed"] else "No", s["label"], s["status"], dec_text(s), dec_ids(s),
                        "; ".join(s["a_chunks"]), "; ".join(s["b_chunks"])])
    with (args.out / "edc_only_evaluation.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Source ID", "Drug name", "Original source sentence", "Chunk ID", "Triple ID", "Subject", "Relation",
                    "Object", "From this source? (Yes/No)", "Correct/Incorrect", "Comment"])
        for r in a_rows:
            w.writerow([r["source_id"], r["drug"], r["original"], r["chunk_id"], r["triple_id"], r["subject"],
                        r["relation"], r["object"], "", "", ""])
    with (args.out / "preprocessing_edc_evaluation.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Source ID", "Drug name", "Original source sentence", "Coreference-resolved sentence",
                    "Decomposed sentence", "Decomposed sentence ID", "Chunk ID", "Triple ID", "Subject", "Relation",
                    "Object", "From this source? (Yes/No)", "Correct/Incorrect", "Comment"])
        for r in b_rows:
            w.writerow([r["source_id"], r["drug"], r["original"], r["coref"], dec_text(r), dec_ids(r), r["chunk_id"],
                        r["triple_id"], r["subject"], r["relation"], r["object"], "", "", ""])
    used_a = sorted({c for s in sampled for c in s["a_chunks"]})
    used_b = sorted({c for s in sampled for c in s["b_chunks"]})
    with (args.out / "chunk_texts.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Model", "Chunk ID", "Chunk text given to EDC"])
        for c in used_a:
            w.writerow(["EDC only", c, a_chunk_text.get(c, "")])
        for c in used_b:
            w.writerow(["Preprocessing + EDC", c, b_chunk_text.get(c, "")])

    build_xlsx(args.out / "comparison.xlsx", sampled, a_rows, b_rows, used_a, used_b, a_chunk_text, b_chunk_text,
               dec_text)

    # ---- summary
    def per_source(rows):
        c = collections.Counter(r["source_id"] for r in rows if r["triple_id"])
        return c
    ca, cb = per_source(a_rows), per_source(b_rows)
    cb_first = collections.Counter(r["source_id"] for r in b_rows if r["triple_id"] and r.get("first_provenance"))
    summary = {
        "drugs_sampled": len({s["source_file"] for s in sampled}), "sources_sampled": len(sampled),
        "linkable_sources_total": sum(len(v) for v in universe.values()), "original_sentences_total": len(raw),
        "excluded_from_sampling_pool": dict(reasons),
        "edc_only_triples": sum(ca.values()), "preprocessing_edc_triples": sum(cb.values()),
        "preprocessing_edc_triples_first_provenance_only": sum(cb_first.values()),
        "edc_only_triples_per_source": round(sum(ca.values()) / len(sampled), 2),
        "preprocessing_edc_triples_per_source": round(sum(cb.values()) / len(sampled), 2),
        "sources_with_no_triple_edc_only": sum(1 for s in sampled if ca[s["source_id"]] == 0),
        "sources_with_no_triple_preprocessing_edc": sum(1 for s in sampled if cb[s["source_id"]] == 0),
        "sampled_sources_with_mapping_failure": 0,
        "a_chunk_rebuild_mismatches": len(mismatch),
        "note": "triples are chunk-level candidates: every triple of the chunk containing the source sentence",
    }
    (args.out / "summary.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(summary, indent=1, ensure_ascii=False), flush=True)
    return 0


def build_xlsx(path, sampled, a_rows, b_rows, used_a, used_b, a_text, b_text, dec_text):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    lines = [
        "Human evaluation: EDC only (Model A) vs pronoun resolution + sentence decomposition + EDC (Model B)",
        "",
        "Sheet 'Comparison': one block per original source sentence (150 sentences, 3 per drug, random seed 42).",
        "Each block lists every triple both models extracted from the chunk that contains the sentence.",
        "EDC extracts from chunks of about 8 sentences and records only the chunk, so a block also contains triples",
        "that came from the chunk's OTHER sentences. Nothing was matched by similarity.",
        "",
        "How to evaluate each triple row:",
        "  1. From this source?  Yes = the triple is stated in the original sentence of this block; No = it comes from",
        "     another sentence of the chunk (see sheet 'Chunks' for the full chunk text).",
        "  2. Correct/Incorrect (only for Yes rows): is the triple a correct statement of the original sentence?",
        "  3. Comment: free text (wrong relation, missing qualifier, wrong entity, ...).",
        "",
        "Rows saying '(no triple extracted from this chunk)' mean the model produced no triple for that chunk.",
        "Model A triples come from the EDC-only explorer data, which keeps only the first chunk of a triple seen in",
        "several chunks; such a triple can be missing from a later chunk's block (about 5% of Model A triples).",
        "Model B: the decomposed sentences listed are those of THIS source sentence; EDC read them together with the",
        "rest of the chunk, and which decomposed sentence produced a triple is not recorded.",
    ]
    for i, l in enumerate(lines, 1):
        readme.cell(row=i, column=1, value=l)
    readme.column_dimensions["A"].width = 120
    readme["A1"].font = Font(bold=True, size=13)

    ws = wb.create_sheet("Comparison")
    head = ["No", "Source ID", "Drug", "Model", "Chunk ID", "Triple ID", "Subject", "Relation", "Object",
            "From this source? (Yes/No)", "Correct/Incorrect", "Comment"]
    ws.append(head)
    fill_src = PatternFill("solid", fgColor="E7E6E6")
    fill_a = PatternFill("solid", fgColor="DDEBF7")
    fill_b = PatternFill("solid", fgColor="E2EFDA")
    fill_head = PatternFill("solid", fgColor="404040")
    thin = Side(style="thin", color="BFBFBF")
    wrap = Alignment(wrap_text=True, vertical="top")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill_head
        c.alignment = wrap
    dv_yes = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    dv_ok = DataValidation(type="list", formula1='"Correct,Incorrect"', allow_blank=True)
    ws.add_data_validation(dv_yes)
    ws.add_data_validation(dv_ok)
    by_a, by_b = collections.defaultdict(list), collections.defaultdict(list)
    for r in a_rows:
        by_a[r["source_id"]].append(r)
    for r in b_rows:
        by_b[r["source_id"]].append(r)
    for s in sampled:
        info = [("Original sentence", s["original"]),
                ("Pronoun-resolved (B)", s["coref"] + ("" if s["coref_changed"] else "   [unchanged]")),
                ("Decomposed (B)", f"[{s['label']} -> {s['status']}]\n" + dec_text(s))]
        for k, (name, val) in enumerate(info):
            ws.append([s["no"] if k == 0 else "", s["source_id"] if k == 0 else "", s["drug"] if k == 0 else "",
                       name, "", "", val])
            row = ws.max_row
            ws.merge_cells(start_row=row, start_column=7, end_row=row, end_column=12)
            for c in ws[row]:
                c.fill = fill_src
                c.alignment = wrap
            ws.cell(row=row, column=4).font = Font(bold=True)
            if k == 0:
                ws.cell(row=row, column=7).font = Font(bold=True)
            ws.row_dimensions[row].height = max(18, 15 * (1 + len(val) // 110 + val.count("\n")))
        for model, rows, fill in (("A: EDC only", by_a[s["source_id"]], fill_a),
                                  ("B: Preprocessing + EDC", by_b[s["source_id"]], fill_b)):
            for r in rows:
                ws.append([s["no"], s["source_id"], s["drug"], model, r["chunk_id"], r["triple_id"], r["subject"],
                           r["relation"], r["object"], "", "", ""])
                row = ws.max_row
                for c in ws[row]:
                    c.fill = fill
                    c.alignment = wrap
                    c.border = Border(bottom=thin)
                if r["triple_id"]:
                    dv_yes.add(ws.cell(row=row, column=10))
                    dv_ok.add(ws.cell(row=row, column=11))
                else:
                    ws.cell(row=row, column=9).font = Font(italic=True, color="7F7F7F")
        ws.append([])
    for col, width in zip("ABCDEFGHIJKL", [5, 24, 14, 22, 30, 9, 28, 24, 40, 14, 14, 30]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"

    for title, ids, texts in (("Chunks A (EDC only)", used_a, a_text), ("Chunks B (Preproc+EDC)", used_b, b_text)):
        cs = wb.create_sheet(title)
        cs.append(["Chunk ID", "Chunk text given to EDC"])
        for c in cs[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = fill_head
        for cid in ids:
            cs.append([cid, texts.get(cid, "")])
            cs.cell(row=cs.max_row, column=2).alignment = wrap
        cs.column_dimensions["A"].width = 42
        cs.column_dimensions["B"].width = 120

    ss = wb.create_sheet("Sources")
    ss.append(["No", "Source ID", "Drug", "Original source sentence", "Pronoun-resolved sentence", "Decomposed sentences",
               "A chunk", "B chunk(s)", "A triples in chunk", "B triples in chunk(s)"])
    for c in ss[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill_head
    for s in sampled:
        ss.append([s["no"], s["source_id"], s["drug"], s["original"], s["coref"], dec_text(s), "; ".join(s["a_chunks"]),
                   "; ".join(s["b_chunks"]), sum(1 for r in by_a[s["source_id"]] if r["triple_id"]),
                   sum(1 for r in by_b[s["source_id"]] if r["triple_id"])])
        for c in ss[ss.max_row]:
            c.alignment = wrap
    for col, width in zip("ABCDEFGHIJ", [5, 24, 14, 60, 60, 60, 30, 30, 10, 10]):
        ss.column_dimensions[col].width = width
    ss.freeze_panes = "A2"
    wb.save(path)


if __name__ == "__main__":
    raise SystemExit(main())
