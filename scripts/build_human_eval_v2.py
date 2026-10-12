#!/usr/bin/env python3
"""Side-by-side human-evaluation sheet: original sentence | Model A triples | Model B triples.

Model A = EDC only (text_edc), Model B = pronoun resolution + decomposition + EDC (coref_decomp_edc_r2).
Existing triples only; nothing is re-extracted or edited.

EDC records only the chunk (about 8 sentences) a triple came from, so each triple of a sampled sentence's chunk is
assigned to a sentence of that chunk by content, with the method recorded per triple:
  exact       the triple's object appears word-for-word in the sentence
  word        otherwise, >= 60% of the object's content words (5-letter stems) are in the sentence and that sentence is
              the single best match in the chunk
  (unassigned triples are counted, not shown)
Model A sentences = original sentences of the chunk; Model B sentences = the decomposed sentences of the chunk.

Sampling: per drug, original sentences shuffled with random.Random(42); drugs are visited round-robin, one sentence
at a time, until both models have at least --target assigned triples. Sentences with no triple are kept.
"""

from __future__ import annotations

import argparse
import collections
import csv
import json
import random
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from build_human_eval import a_line_chunks, norm, source_id  # noqa: E402
from classify_drugs_pilot_sentences import ensure_nltk, extract_directory_records  # noqa: E402

EVALUATORS = 2                     # independent evaluators; each gets its own judgement + comment columns
SOURCE_URL: dict[str, str] = {}   # page file -> Drugs.com URL (filled in main)
STOP = {"the", "and", "for", "with", "you", "your", "are", "may", "can", "not", "this", "that", "has", "have", "from",
        "into", "use", "used", "using", "take", "taking", "any", "all", "other", "medicine", "medicines", "such",
        "will", "been", "being", "its", "their", "them", "they", "who", "which", "when", "than", "more", "also"}


def key(text: str) -> str:
    return " " + re.sub(r"[^a-z0-9]+", " ", text.lower()).strip() + " "


def words(text: str) -> set[str]:
    return {w[:5] for w in re.findall(r"[a-z0-9]+", text.lower()) if len(w) > 2 and w not in STOP}


def assign(obj: str, subj: str, sentences: list[str]) -> tuple[list[int], str]:
    """Indices of the sentence(s) a triple belongs to, and how it was decided."""
    o = key(obj)
    exact = [i for i, s in enumerate(sentences) if o.strip() and o in key(s)]
    if exact:
        return exact, "exact"
    ow = words(obj)
    if not ow:
        return [], "none"
    scores = []
    for i, s in enumerate(sentences):
        sw = words(s)
        scores.append((len(ow & sw) / len(ow), len(words(subj) & sw), i))
    scores.sort(reverse=True)
    best = scores[0]
    tie = len(scores) > 1 and scores[1][:2] == best[:2]
    if best[0] >= 0.6 and not tie:
        return [best[2]], "word"
    return [], "none"


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--raw-dir", type=Path, default=ROOT / "data/text-only/text-only")
    ap.add_argument("--a-html", type=Path, default=ROOT / "drug-kg/output/viz/text_edc/kg_explorer.html")
    ap.add_argument("--b-exp", type=Path, default=ROOT / "drug-kg/output/experiments/coref_decomp_edc_r2")
    ap.add_argument("--trace", type=Path, default=ROOT / "data/text-only/coref_pipeline/4_trace/trace.jsonl")
    ap.add_argument("--out", type=Path, default=ROOT / "data/human_eval_v2")
    ap.add_argument("--target", type=int, default=150)
    ap.add_argument("--seed", type=int, default=42)
    args = ap.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)

    ensure_nltk()
    raw, _ = extract_directory_records(args.raw_dir)

    # Model A: chunk -> triples, chunk -> original sentences (via the chunker's own line grouping)
    html = args.a_html.read_text(encoding="utf-8")
    A = json.loads(re.search(r'<script\b[^>]*\bid=["\']kg-data["\'][^>]*>(.*?)</script>', html, re.S).group(1))
    lab = {n["id"]: n["l"] for n in A["nodes"]}
    a_trip = collections.defaultdict(list)
    for i, e in enumerate(A["edges"]):
        if "edc" in e["src"] and e.get("ch"):
            a_trip[e["ch"]].append({"id": f"A{i:05d}", "s": lab[e["s"]], "r": e["r"], "o": lab[e["o"]]})
    a_lines = {p.name: a_line_chunks(p) for p in args.raw_dir.glob("*.txt")}
    a_chunk_sents = collections.defaultdict(list)   # chunk -> [(source key, sentence)]
    src_a_chunk = {}
    for r in raw:
        k = (r["source_file"], r["line_number"], r["sentence_index"])
        for cid, piece in a_lines[r["source_file"]].get(r["line_number"], []):
            if norm(r["sentence"]) in norm(piece[2:] if piece.startswith("- ") else piece):
                a_chunk_sents[cid].append((k, r["sentence"]))
                src_a_chunk[k] = cid
                break

    # Model B: chunk -> triples, chunk -> decomposed sentences (with their original-sentence record id)
    trace = {}
    for line in args.trace.open(encoding="utf-8"):
        t = json.loads(line)
        f, ln, si, _ = t["record_id"].split("|")
        trace[(f, int(ln), int(si))] = t
    rec2key = {t["record_id"]: k for k, t in trace.items()}
    b_chunk_sents = collections.defaultdict(list)    # chunk -> [(source key, output_index, sentence)]
    src_b_chunks = collections.defaultdict(set)
    for path in (args.b_exp / "chunks").glob("*.jsonl"):
        for line in path.open(encoding="utf-8"):
            c = json.loads(line)
            for s in c["sentences"]:
                k = rec2key.get(s["record_id"])
                b_chunk_sents[c["chunk_id"]].append((k, s["output_index"], s["sentence"]))
                if k:
                    src_b_chunks[k].add(c["chunk_id"])
    b_trip = collections.defaultdict(list)
    for n, line in enumerate((args.b_exp / "kg/triples.jsonl").open(encoding="utf-8")):
        t = json.loads(line)
        for cid in dict.fromkeys(p["chunk_id"] for p in t.get("provenance", []) if p.get("source") == "edc"):
            b_trip[cid].append({"id": f"B{n:05d}", "s": t["subject"], "r": t["relation"], "o": t["object"]})

    # assign every triple of every chunk to sentence(s)
    a_by_src, b_by_src = collections.defaultdict(list), collections.defaultdict(list)
    stats = collections.Counter()
    for cid, sents in a_chunk_sents.items():
        texts = [s for _, s in sents]
        for t in a_trip.get(cid, []):
            idx, how = assign(t["o"], t["s"], texts)
            stats[f"A_{how}"] += 1
            for i in idx:
                a_by_src[sents[i][0]].append({**t, "how": how, "chunk": cid})
    for cid, sents in b_chunk_sents.items():
        texts = [s for _, _, s in sents]
        for t in b_trip.get(cid, []):
            idx, how = assign(t["o"], t["s"], texts)
            stats[f"B_{how}"] += 1
            for i in idx:
                k, oi, sent = sents[i]
                if k:
                    b_by_src[k].append({**t, "how": how, "chunk": cid, "dec_id": f"D{oi + 1}", "dec": sent})

    # sampling pool: original sentences linked to a chunk in both models
    pool = collections.defaultdict(list)
    for r in raw:
        k = (r["source_file"], r["line_number"], r["sentence_index"])
        t = trace.get(k)
        if k in src_a_chunk and t and t["raw_match"] == "sentence" and src_b_chunks.get(k):
            pool[r["source_file"]].append((k, r))
    rng = random.Random(args.seed)
    files = sorted(pool)
    for f in files:
        rng.shuffle(pool[f])
    picked, na, nb, round_ = [], 0, 0, 0
    while na < args.target or nb < args.target:
        progressed = False
        for f in files:
            if round_ < len(pool[f]) and (na < args.target or nb < args.target):
                k, r = pool[f][round_]
                picked.append((k, r))
                na += len({t["id"] for t in a_by_src[k]})
                nb += len({t["id"] for t in b_by_src[k]})
                progressed = True
        if not progressed:
            break
        round_ += 1

    def dedupe(ts):
        seen, out = set(), []
        for t in ts:
            if t["id"] not in seen:
                seen.add(t["id"])
                out.append(t)
        return out

    rows = []
    for no, (k, r) in enumerate(picked, 1):
        t = trace[k]
        sec = r["section_title"] or ""
        if r.get("subsection_title"):
            sec += f" > {r['subsection_title']}"
        rows.append({"no": no, "sid": source_id(*k), "drug": r["drug"], "orig": r["sentence"], "section": sec,
                     "doc": f"{r['source_file']} ({SOURCE_URL.get(r['source_file'], '')})",
                     "coref": t["coref_sentence"], "coref_changed": bool(t["coref_edits"]), "status": t["status"],
                     "dec": [(f"D{d['output_index'] + 1}", d["sentence"]) for d in t["decomposed"]],
                     "A": dedupe(a_by_src[k]), "B": dedupe(b_by_src[k]),
                     "a_chunk": src_a_chunk[k], "b_chunks": sorted(src_b_chunks[k])})

    for p in args.raw_dir.glob("*.txt"):
        m = re.search(r"^Source: (\S+)", p.read_text(encoding="utf-8"), re.M)
        SOURCE_URL[p.name] = m.group(1) if m else ""
    for r in rows:
        r["doc"] = f"{r['doc'].split(' (')[0]} ({SOURCE_URL.get(r['doc'].split(' (')[0], '')})"
    write_outputs(args.out, rows)
    s = {"sentences": len(rows), "drugs": len({r["drug"] for r in rows}),
         "triples_A": sum(len(r["A"]) for r in rows), "triples_B": sum(len(r["B"]) for r in rows),
         "per_sentence_A": round(sum(len(r["A"]) for r in rows) / len(rows), 2),
         "per_sentence_B": round(sum(len(r["B"]) for r in rows) / len(rows), 2),
         "sentences_no_triple_A": sum(not r["A"] for r in rows), "sentences_no_triple_B": sum(not r["B"] for r in rows),
         "sentences_no_triple_both": sum(not r["A"] and not r["B"] for r in rows),
         "match_method_A": dict(collections.Counter(t["how"] for r in rows for t in r["A"])),
         "match_method_B": dict(collections.Counter(t["how"] for r in rows for t in r["B"])),
         "assignment_over_all_chunks": dict(stats)}
    (args.out / "summary.json").write_text(json.dumps(s, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(s, indent=1, ensure_ascii=False))
    return 0


def write_outputs(out: Path, rows: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    def fmt(t):
        return f"({t['s']}, {t['r']}, {t['o']})"

    # flat CSVs (one row per triple; a sentence with no triple keeps one row saying so)
    with (out / "sampled_sources.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["No", "Source ID", "Drug name", "Source document", "Section title", "Original source sentence",
                    "Coreference-resolved sentence",
                    "Decomposed sentences", "Model A triples", "Model B triples"])
        for r in rows:
            w.writerow([r["no"], r["sid"], r["drug"], r["doc"], r["section"], r["orig"], r["coref"],
                        "\n".join(f"[{i}] {s}" for i, s in r["dec"]), len(r["A"]), len(r["B"])])
    for name, model in (("edc_only_evaluation.csv", "A"), ("preprocessing_edc_evaluation.csv", "B")):
        with (out / name).open("w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            head = ["Source ID", "Drug name", "Source document", "Section title", "Original source sentence"]
            if model == "B":
                head += ["Coreference-resolved sentence", "Decomposed sentence", "Decomposed sentence ID"]
            evals = [c for e in range(1, EVALUATORS + 1) for c in (f"Evaluator {e} Correct/Incorrect", f"Evaluator {e} Comment")]
            w.writerow(head + ["Triple ID", "Subject", "Relation", "Object", "Match method", "Chunk ID"] + evals)
            for r in rows:
                base = [r["sid"], r["drug"], r["doc"], r["section"], r["orig"]]
                ts = r[model] or [None]
                for t in ts:
                    extra = [r["coref"], t["dec"] if t else "", f"{r['sid']}:{t['dec_id']}" if t else ""] if model == "B" else []
                    if t:
                        w.writerow(base + extra + [t["id"], t["s"], t["r"], t["o"], t["how"], t["chunk"]] + [""] * len(evals))
                    else:
                        w.writerow(base + extra + ["", "", "", "(no triple extracted)", "", ""] + [""] * len(evals))

    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison"
    # columns: (header, width, group) — group decides the fill; the evaluator columns repeat per evaluator
    src_cols = [("No", 5), ("Source ID", 22), ("Drug", 13), ("Source document", 26), ("Section title", 24),
                ("Original sentence", 48), ("Decomposed sentences (B)", 48)]
    def model_cols(m, extra):
        cols = [(f"Model {m} triple", 42)] + extra + [(f"{m} match", 8)]
        for e in range(1, EVALUATORS + 1):
            cols += [(f"{m} Evaluator {e} Correct/Incorrect", 13), (f"{m} Evaluator {e} Comment", 20)]
        return cols
    a_cols, b_cols = model_cols("A", []), model_cols("B", [("B from", 7)])
    cols = src_cols + a_cols + b_cols
    ws.append([h for h, _ in cols])
    pos = {h: i + 1 for i, (h, _) in enumerate(cols)}
    n_src, n_a = len(src_cols), len(a_cols)
    dark, fa, fb, fs = (PatternFill("solid", fgColor=c) for c in ("404040", "DDEBF7", "E2EFDA", "F2F2F2"))
    fa2, fb2 = PatternFill("solid", fgColor="C5D9F1"), PatternFill("solid", fgColor="C6E0B4")  # evaluator 2: darker
    wrap = Alignment(wrap_text=True, vertical="top")
    sep = Side(style="medium", color="808080")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = dark
        c.alignment = wrap
    dv = DataValidation(type="list", formula1='"Correct,Incorrect"', allow_blank=True)
    ws.add_data_validation(dv)
    for r in rows:
        n = max(len(r["A"]), len(r["B"]), 1)
        start = ws.max_row + 1
        dec = "\n".join(f"[{i}] {s}" for i, s in r["dec"])
        if r["coref_changed"]:
            dec = f"(pronoun resolved) {r['coref']}\n" + dec
        for j in range(n):
            a = r["A"][j] if j < len(r["A"]) else None
            b = r["B"][j] if j < len(r["B"]) else None
            vals = {"No": r["no"], "Source ID": r["sid"], "Drug": r["drug"], "Source document": r["doc"],
                    "Section title": r["section"], "Original sentence": r["orig"], "Decomposed sentences (B)": dec,
                    "Model A triple": fmt(a) if a else ("(no triple extracted)" if j == 0 and not r["A"] else ""),
                    "A match": a["how"] if a else "",
                    "Model B triple": fmt(b) if b else ("(no triple extracted)" if j == 0 and not r["B"] else ""),
                    "B from": b["dec_id"] if b else "", "B match": b["how"] if b else ""}
            ws.append([vals.get(h, "") for h, _ in cols])
            row = ws.max_row
            for col in range(1, len(cols) + 1):
                cell = ws.cell(row=row, column=col)
                cell.alignment = wrap
                head = cols[col - 1][0]
                if col <= n_src:
                    cell.fill = fs
                elif col <= n_src + n_a:
                    cell.fill = fa2 if "Evaluator 2" in head else fa
                else:
                    cell.fill = fb2 if "Evaluator 2" in head else fb
            for m, t in (("A", a), ("B", b)):
                if t:
                    for e in range(1, EVALUATORS + 1):
                        dv.add(ws.cell(row=row, column=pos[f"{m} Evaluator {e} Correct/Incorrect"]))
                elif j == 0:
                    ws.cell(row=row, column=pos[f"Model {m} triple"]).font = Font(italic=True, color="7F7F7F")
        for col in range(1, n_src + 1):
            if n > 1:
                ws.merge_cells(start_row=start, start_column=col, end_row=start + n - 1, end_column=col)
        for col in range(1, len(cols) + 1):
            ws.cell(row=start, column=col).border = Border(top=sep)
    from openpyxl.utils import get_column_letter
    for i, (_, width) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "D2"   # header row + No/Source ID/Drug; wider frozen areas block horizontal scrolling

    guide = wb.create_sheet("How to evaluate")
    for i, line in enumerate([
        "One block per original sentence. Left: the original sentence (and how Model B decomposed it).",
        "Middle (blue): triples Model A (EDC only) extracted for this sentence. Right (green): triples Model B",
        "(pronoun resolution + decomposition + EDC) extracted for it. Rows are NOT matched pairs; each side is just a list.",
        "",
        "Two evaluators judge independently: Evaluator 1 fills only the 'Evaluator 1' columns, Evaluator 2 only the",
        "'Evaluator 2' columns (darker shade), without looking at each other's answers.",
        "For every triple choose Correct or Incorrect: is it a correct statement of the original sentence?",
        "'(no triple extracted)' = the model produced no triple for this sentence. Leave it, or note in Comment if",
        "the sentence clearly contains a fact that should have been extracted.",
        "",
        "Match = how the triple was linked to this sentence (EDC records only the ~8-sentence chunk, not the sentence):",
        "  exact = the triple's object appears word for word in the sentence",
        "  word  = most of the object's words appear in the sentence, and it is the best match in the chunk",
        "If a triple clearly belongs to a different sentence, mark it Incorrect and write 'wrong sentence' in Comment.",
        "B from = which decomposed sentence (D1, D2, ...) the Model B triple was linked to.",
    ], 1):
        guide.cell(row=i, column=1, value=line)
    guide.column_dimensions["A"].width = 120
    wb.move_sheet("How to evaluate", offset=-1)
    wb.save(out / "comparison.xlsx")


if __name__ == "__main__":
    raise SystemExit(main())
