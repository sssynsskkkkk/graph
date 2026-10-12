#!/usr/bin/env python3
"""Whole-page human-evaluation sheet for ONE drug: every original sentence of the page in page order, with the
triples Model A (EDC only, text_edc) and Model B (pronoun resolution + decomposition + EDC, coref_decomp_edc_r2)
extracted for it, side by side. Existing triples only (EDC-extracted; page-metadata triples excluded).

Drug choice: drugs whose mean EDC triple count over the two models is within --target +/- --tolerance;
one is drawn with random.Random(42).
Triple -> sentence linking is by content within the triple's chunk (EDC records only the chunk), method per triple:
exact (object word for word in the sentence) or word (>= 60% of the object's content words, single best sentence).
Triples that cannot be linked are listed at the end so that every triple of the drug appears in the sheet.
"""

from __future__ import annotations

import argparse
import collections
import csv
import json
import random
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from build_human_eval import a_line_chunks, norm  # noqa: E402
from build_human_eval_v2 import assign  # noqa: E402
from classify_drugs_pilot_sentences import ensure_nltk, extract_directory_records  # noqa: E402


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--raw-dir", type=Path, default=ROOT / "data/text-only/text-only")
    ap.add_argument("--a-html", type=Path, default=ROOT / "drug-kg/output/viz/text_edc/kg_explorer.html")
    ap.add_argument("--b-exp", type=Path, default=ROOT / "drug-kg/output/experiments/coref_decomp_edc_r2")
    ap.add_argument("--trace", type=Path, default=ROOT / "data/text-only/coref_pipeline/4_trace/trace.jsonl")
    ap.add_argument("--out", type=Path, default=ROOT / "data/human_eval_v3")
    ap.add_argument("--target", type=int, default=150)
    ap.add_argument("--tolerance", type=int, default=20)
    ap.add_argument("--seed", type=int, default=42)
    ap.add_argument("--drug", default=None, help="force a drug file stem instead of sampling")
    args = ap.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)

    # ---- triples of both models, keyed by chunk
    html = args.a_html.read_text(encoding="utf-8")
    A = json.loads(re.search(r'<script\b[^>]*\bid=["\']kg-data["\'][^>]*>(.*?)</script>', html, re.S).group(1))
    lab = {n["id"]: n["l"] for n in A["nodes"]}
    a_trip = collections.defaultdict(list)
    for i, e in enumerate(A["edges"]):
        if "edc" in e["src"] and e.get("ch"):
            a_trip[e["ch"]].append({"id": f"A{i:05d}", "s": lab[e["s"]], "r": e["r"], "o": lab[e["o"]]})
    b_trip = collections.defaultdict(list)
    for n, line in enumerate((args.b_exp / "kg/triples.jsonl").open(encoding="utf-8")):
        t = json.loads(line)
        for cid in dict.fromkeys(p["chunk_id"] for p in t.get("provenance", []) if p.get("source") == "edc"):
            b_trip[cid].append({"id": f"B{n:05d}", "s": t["subject"], "r": t["relation"], "o": t["object"]})

    def per_drug(trip):
        c = collections.defaultdict(set)
        for cid, ts in trip.items():
            c[cid.split("#")[0]].update(t["id"] for t in ts)
        return {k: len(v) for k, v in c.items()}
    ca, cb = per_drug(a_trip), per_drug(b_trip)
    cands = sorted(d for d in ca if abs((ca[d] + cb.get(d, 0)) / 2 - args.target) <= args.tolerance)
    drug = args.drug or random.Random(args.seed).choice(cands)
    fname = f"{drug}.txt"

    # ---- original sentences of the page and their chunks
    ensure_nltk()
    raw = [r for r in extract_directory_records(args.raw_dir)[0] if r["source_file"] == fname]
    lines = a_line_chunks(args.raw_dir / fname)
    trace = {}
    for line in args.trace.open(encoding="utf-8"):
        t = json.loads(line)
        f, ln, si, _ = t["record_id"].split("|")
        if f == fname:
            trace[(int(ln), int(si))] = t
    rec2key = {t["record_id"]: k for k, t in trace.items()}

    sents = []      # page order
    a_chunk_sents = collections.defaultdict(list)
    for r in raw:
        k = (r["line_number"], r["sentence_index"])
        a_ch = next((cid for cid, piece in lines.get(r["line_number"], [])
                     if norm(r["sentence"]) in norm(piece[2:] if piece.startswith("- ") else piece)), None)
        t = trace.get(k)
        sents.append({"k": k, "section": r["section_title"] or "", "orig": r["sentence"], "a_chunk": a_ch,
                      "coref": t["coref_sentence"] if t else r["sentence"], "coref_changed": bool(t and t["coref_edits"]),
                      "status": t["status"] if t else "", "label": t["label"] if t else "",
                      "dec": [(f"D{d['output_index'] + 1}", d["sentence"]) for d in t["decomposed"]] if t else [],
                      "A": [], "B": []})
        if a_ch:
            a_chunk_sents[a_ch].append(len(sents) - 1)
    b_chunk_sents = collections.defaultdict(list)   # chunk -> [(sentence index, D id, decomposed sentence)]
    b_chunks_of = collections.defaultdict(set)
    idx_of = {s["k"]: i for i, s in enumerate(sents)}
    for line in (args.b_exp / "chunks" / f"{drug}.jsonl").open(encoding="utf-8"):
        c = json.loads(line)
        for s in c["sentences"]:
            k = rec2key.get(s["record_id"])
            i = idx_of.get(k) if k else None
            b_chunk_sents[c["chunk_id"]].append((i, f"D{s['output_index'] + 1}", s["sentence"]))
            if i is not None:
                b_chunks_of[i].add(c["chunk_id"])

    # ---- link triples to sentences
    unlinked = {"A": [], "B": []}
    for cid, ts in a_trip.items():
        if cid.split("#")[0] != drug:
            continue
        idxs = a_chunk_sents.get(cid, [])
        for t in ts:
            hit, how = assign(t["o"], t["s"], [sents[i]["orig"] for i in idxs]) if idxs else ([], "none")
            for h in hit:
                sents[idxs[h]]["A"].append({**t, "how": how, "chunk": cid})
            if not hit:
                unlinked["A"].append({**t, "how": "-", "chunk": cid})
    for cid, ts in b_trip.items():
        if cid.split("#")[0] != drug:
            continue
        items = b_chunk_sents.get(cid, [])
        for t in ts:
            hit, how = assign(t["o"], t["s"], [x[2] for x in items]) if items else ([], "none")
            placed = False
            for h in hit:
                i, did, dsent = items[h]
                if i is not None:
                    sents[i]["B"].append({**t, "how": how, "chunk": cid, "dec_id": did, "dec": dsent})
                    placed = True
            if not placed:
                unlinked["B"].append({**t, "how": "-", "chunk": cid, "dec_id": "", "dec": ""})
    for s in sents:
        for m in ("A", "B"):
            seen, keep = set(), []
            for t in s[m]:
                if t["id"] not in seen:
                    seen.add(t["id"])
                    keep.append(t)
            s[m] = keep

    write(args.out, drug, sents, unlinked)
    uniq = lambda m: len({t["id"] for s in sents for t in s[m]} | {t["id"] for t in unlinked[m]})
    summary = {"candidates": {d: [ca[d], cb.get(d, 0)] for d in cands}, "drug": drug,
               "page_sentences": len(sents), "sentences_given_to_edc": sum(1 for s in sents if s["a_chunk"]),
               "triples_A": uniq("A"), "triples_B": uniq("B"),
               "linked_A": len({t["id"] for s in sents for t in s["A"]}), "linked_B": len({t["id"] for s in sents for t in s["B"]}),
               "unlinked_A": len(unlinked["A"]), "unlinked_B": len(unlinked["B"]),
               "sentences_no_triple_A": sum(1 for s in sents if s["a_chunk"] and not s["A"]),
               "sentences_no_triple_B": sum(1 for s in sents if s["a_chunk"] and not s["B"]),
               "match_A": dict(collections.Counter(t["how"] for s in sents for t in s["A"])),
               "match_B": dict(collections.Counter(t["how"] for s in sents for t in s["B"]))}
    (args.out / "summary.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(summary, indent=1, ensure_ascii=False))
    return 0


def write(out: Path, drug: str, sents: list[dict], unlinked: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    fmt = lambda t: f"({t['s']}, {t['r']}, {t['o']})"
    wb = Workbook()
    guide = wb.active
    guide.title = "How to evaluate"
    for i, line in enumerate([
        f"Drug: {drug}. Every sentence of the Drugs.com page, in page order (sheet 'Comparison').",
        "Left (grey): the original sentence and how Model B decomposed it. Middle (blue): triples Model A (EDC only)",
        "extracted for it. Right (green): triples Model B (pronoun resolution + decomposition + EDC) extracted for it.",
        "Rows are NOT matched pairs; each side is a list. '(no triple extracted)' = that model has no triple here.",
        "Grey italic 'not given to EDC' = section excluded from extraction (e.g. Popular FAQ) for both models.",
        "",
        "For every triple choose Correct or Incorrect: is it a correct statement of the page? Add a Comment if useful",
        "(wrong relation, missing condition, wrong entity, belongs to another sentence, ...).",
        "",
        "Match = how a triple was linked to the sentence (EDC records only the ~8-sentence chunk it read):",
        "  exact = the triple's object appears word for word in the sentence",
        "  word  = most of the object's words appear in the sentence, and it is the best match in the chunk",
        "The last block 'NOT LINKED TO A SENTENCE' holds triples that could not be linked; they are still the",
        "model's output for this page and should be evaluated too.",
    ], 1):
        guide.cell(row=i, column=1, value=line)
    guide.column_dimensions["A"].width = 120
    guide["A1"].font = Font(bold=True, size=13)

    ws = wb.create_sheet("Comparison")
    ws.append(["No", "Section", "Original sentence", "Decomposed sentences (B)",
               "Model A triple", "A match", "A Correct/Incorrect", "A Comment",
               "Model B triple", "B from", "B match", "B Correct/Incorrect", "B Comment"])
    dark, fa, fb, fs, fsec = (PatternFill("solid", fgColor=c) for c in ("404040", "DDEBF7", "E2EFDA", "F2F2F2", "D9D9D9"))
    wrap = Alignment(wrap_text=True, vertical="top")
    sep = Side(style="thin", color="808080")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = dark
        c.alignment = wrap
    dv = DataValidation(type="list", formula1='"Correct,Incorrect"', allow_blank=True)
    ws.add_data_validation(dv)

    def block(no, section, orig, dec, A, B, given=True):
        n = max(len(A), len(B), 1)
        start = ws.max_row + 1
        for j in range(n):
            a = A[j] if j < len(A) else None
            b = B[j] if j < len(B) else None
            none = "(no triple extracted)" if given else "(not given to EDC)"
            ws.append([no, section, orig, dec,
                       fmt(a) if a else (none if j == 0 and not A else ""), a["how"] if a else "", "", "",
                       fmt(b) if b else (none if j == 0 and not B else ""), b.get("dec_id", "") if b else "",
                       b["how"] if b else "", "", ""])
            row = ws.max_row
            for col in range(1, 14):
                cell = ws.cell(row=row, column=col)
                cell.alignment = wrap
                cell.fill = fs if col <= 4 else (fa if col <= 8 else fb)
            if a:
                dv.add(ws.cell(row=row, column=7))
            elif j == 0:
                ws.cell(row=row, column=5).font = Font(italic=True, color="7F7F7F")
            if b:
                dv.add(ws.cell(row=row, column=12))
            elif j == 0:
                ws.cell(row=row, column=9).font = Font(italic=True, color="7F7F7F")
        if n > 1:
            for col in range(1, 5):
                ws.merge_cells(start_row=start, start_column=col, end_row=start + n - 1, end_column=col)
        for col in range(1, 14):
            ws.cell(row=start, column=col).border = Border(top=sep)

    section, no = None, 0
    for s in sents:
        if s["section"] != section:
            section = s["section"]
            ws.append(["", section])
            r = ws.max_row
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=13)
            for col in range(1, 14):
                ws.cell(row=r, column=col).fill = fsec
            ws.cell(row=r, column=2).font = Font(bold=True)
        no += 1
        dec = "\n".join(f"[{i}] {x}" for i, x in s["dec"])
        if s["coref_changed"]:
            dec = f"(pronoun resolved) {s['coref']}\n{dec}"
        block(no, section, s["orig"], dec, s["A"], s["B"], given=bool(s["a_chunk"]))

    ws.append(["", "NOT LINKED TO A SENTENCE (model output for this page whose sentence could not be determined)"])
    r = ws.max_row
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=13)
    for col in range(1, 14):
        ws.cell(row=r, column=col).fill = fsec
    ws.cell(row=r, column=2).font = Font(bold=True)
    block("-", "", "(chunk-level only)", "", unlinked["A"], unlinked["B"])
    for col, width in zip("ABCDEFGHIJKLM", [5, 18, 50, 46, 44, 7, 12, 22, 44, 6, 7, 12, 22]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "B2"   # header row + No; wider frozen areas block horizontal scrolling

    page = wb.create_sheet("Page text")
    page.append(["Original page text (data/text-only/text-only/" + drug + ".txt)"])
    for line in (ROOT / "data/text-only/text-only" / f"{drug}.txt").read_text(encoding="utf-8").splitlines():
        page.append([line])
        page.cell(row=page.max_row, column=1).alignment = Alignment(wrap_text=True)
        if line.startswith("#"):
            page.cell(row=page.max_row, column=1).font = Font(bold=True)
    page.column_dimensions["A"].width = 140
    wb.save(out / f"comparison_{drug}.xlsx")

    # CSV copies (one row per triple)
    for m, name in (("A", "edc_only"), ("B", "preprocessing_edc")):
        with (out / f"{name}_{drug}.csv").open("w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(["No", "Section", "Original sentence"] + (["Decomposed sentence", "Decomposed ID"] if m == "B" else [])
                       + ["Triple ID", "Subject", "Relation", "Object", "Match", "Chunk ID"])
            for no, s in enumerate(sents, 1):
                for t in s[m] or [None]:
                    extra = [t["dec"], t["dec_id"]] if (m == "B" and t) else (["", ""] if m == "B" else [])
                    if t:
                        w.writerow([no, s["section"], s["orig"]] + extra + [t["id"], t["s"], t["r"], t["o"], t["how"], t["chunk"]])
                    else:
                        w.writerow([no, s["section"], s["orig"]] + extra + ["", "", "", "(no triple extracted)" if s["a_chunk"] else "(not given to EDC)", "", ""])
            for t in unlinked[m]:
                w.writerow(["-", "", "(not linked to a sentence)"] + (["", ""] if m == "B" else [])
                           + [t["id"], t["s"], t["r"], t["o"], "-", t["chunk"]])


if __name__ == "__main__":
    raise SystemExit(main())
